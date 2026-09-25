"""Cycle 5 only: linear regression of lab chlorophyll (Chl a, Chl b, Total Chl) on SPAD, for each SPAD source
(whole leaf from SPAD_Data.xlsx; up / mid / low from backup_original/SPAD_Data_B3_dropped_per_part.xlsx).

Writes results/figures/spad_cycle5_regression.svg (PNG to $PREVIEW_DIR if set) and
results/tables/SPAD_Cycle5_Regression.xlsx (fit for all cycle-5 samples and for each block).
"""
import os
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.lines import Line2D
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from scipy import stats

ROOT = Path(__file__).resolve().parent.parent          # project folder
DATA = ROOT / 'data'
FIGS = ROOT / 'results' / 'figures'
TABLES = ROOT / 'results' / 'tables'
CYCLE = 5

# ---- SPAD for cycle 5, per source; part means over the readings present
SPAD = {'Whole leaf': {}}
ws = openpyxl.load_workbook(DATA / 'SPAD_Data.xlsx', data_only=True).worksheets[CYCLE - 1]
for r in range(3, ws.max_row + 1):
    if ws.cell(r, 2).value is not None: SPAD['Whole leaf'][ws.cell(r, 2).value] = ws.cell(r, 3).value
ws = openpyxl.load_workbook(DATA / 'backup_original' / 'SPAD_Data_B3_dropped_per_part.xlsx', data_only=True).worksheets[CYCLE - 1]
for part, label in (('up', 'Upper leaf'), ('mid', 'Mid leaf'), ('low', 'Lower leaf')):
    cols = [j for j in range(3, ws.max_column + 1) if ws.cell(1, j).value == part]
    SPAD[label] = {}
    for r in range(3, ws.max_row + 1):
        v = [ws.cell(r, j).value for j in cols]; v = [x for x in v if isinstance(x, (int, float))]
        if ws.cell(r, 2).value is not None and v: SPAD[label][ws.cell(r, 2).value] = sum(v) / len(v)

lab = [r for r in openpyxl.load_workbook(DATA / 'Chl_Lab_Data.xlsx', data_only=True).active.iter_rows(min_row=2, max_col=5, values_only=True)
       if r[1] is not None and int(r[0]) == CYCLE]
samples = [r[1] for r in lab]
blk = np.array([2 if 'B2' in s else 1 for s in samples])
MEAS = {'Chlorophyll a': np.array([r[2] for r in lab], float), 'Chlorophyll b': np.array([r[3] for r in lab], float),
        'Total chlorophyll': np.array([r[4] for r in lab], float)}

# ---- figure: rows = lab measure, columns = SPAD source
SURF, INK, INK2, GRID = '#fcfcfb', '#0b0b0b', '#3d3c39', '#e2e1dc'
BC = {1: '#2a78d6', 2: '#eb6834'}; BM = {1: 'o', 2: 's'}
BL = {1: 'B1 – outside panel shade', 2: 'B2 – under solar panel'}
plt.rcParams['svg.fonttype'] = 'path'
plt.rcParams.update({'font.family': 'serif', 'font.serif': ['Times New Roman', 'Nimbus Roman', 'STIXGeneral'],
                     'mathtext.fontset': 'stix', 'text.color': INK, 'axes.labelcolor': INK,
                     'xtick.color': INK2, 'ytick.color': INK2, 'font.size': 15})
fig, axes = plt.subplots(3, 4, figsize=(22, 17), facecolor=SURF)
fig.subplots_adjust(left=0.06, right=0.99, top=0.885, bottom=0.09, hspace=0.42, wspace=0.28)
fig.text(0.06, 0.975, f'SPAD vs laboratory chlorophyll, harvest cycle {CYCLE} only', fontsize=28, fontweight='bold', va='top')
fig.text(0.06, 0.94, f'Each point is one sample (n = {len(lab)}: 27 per block). Black line = linear fit to all cycle-{CYCLE} samples. '
         'Columns: SPAD source; rows: laboratory measure.', fontsize=16, color=INK2, va='top')
rec = []
for i, (mname, y) in enumerate(MEAS.items()):
    for j, (src, sp) in enumerate(SPAD.items()):
        ax = axes[i, j]; ax.set_facecolor(SURF)
        x = np.array([sp[s] for s in samples], float)
        for b in (1, 2):
            m = blk == b
            ax.scatter(x[m], y[m], s=58, marker=BM[b], color=BC[b], alpha=0.75, linewidths=0, zorder=3)
        lr = stats.linregress(x, y); xs = np.array([x.min(), x.max()])
        ax.plot(xs, lr.intercept + lr.slope * xs, color=INK, lw=2.4, zorder=4)
        sign = '+' if lr.intercept >= 0 else '−'
        ax.legend(handles=[Line2D([], [], color=BC[b], marker=BM[b], ls='', ms=9, label=BL[b]) for b in (1, 2)] +
                  [Line2D([], [], color=INK, lw=2.4, label=f'Fit: y = {lr.slope:.2f}x {sign} {abs(lr.intercept):.1f}\n'
                                                            f'r = {lr.rvalue:.2f}, R² = {lr.rvalue**2:.2f}')],
                  loc='upper left', frameon=True, facecolor=SURF, edgecolor='none', framealpha=0.9,
                  fontsize=12.5, handletextpad=0.5, borderpad=0.3, labelspacing=0.35)
        ax.set_ylim(0, y.max() * 1.75)
        ax.set_title(f'{mname} – {src} SPAD', loc='left', fontsize=17, fontweight='bold', pad=8)
        ax.set_xlabel(f'{src} SPAD value', fontsize=15, labelpad=5)
        ax.set_ylabel(f'{mname} (lab value)', fontsize=15, labelpad=5)
        ax.grid(color=GRID, lw=1); ax.set_axisbelow(True)
        for s in ('top', 'right'): ax.spines[s].set_visible(False)
        for s in ('left', 'bottom'): ax.spines[s].set_color('#b9b8b2')
        ax.tick_params(length=0, labelsize=13.5)
        for sub, m in (('All cycle-5 samples', np.ones(len(x), bool)), ('B1 outside panel shade', blk == 1), ('B2 under solar panel', blk == 2)):
            l = stats.linregress(x[m], y[m])
            rec.append([src, mname, sub, int(m.sum()), l.slope, l.intercept, l.rvalue, l.rvalue ** 2, l.pvalue])
fig.text(0.06, 0.022, 'Whole leaf = SPAD_Data.xlsx; upper / mid / lower leaf = mean of that part in SPAD_Data_B3_dropped_per_part.xlsx '
         '(readings present only). Font: Nimbus Roman (Times-compatible).', fontsize=14, color=INK2, style='italic')
fig.savefig(FIGS / 'spad_cycle5_regression.svg', facecolor=SURF)
if os.environ.get('PREVIEW_DIR'):
    fig.savefig(Path(os.environ['PREVIEW_DIR']) / 'spad_cycle5_regression.png', dpi=int(os.environ.get('PREVIEW_DPI', 100)), facecolor=SURF)

# ---- table
wb = openpyxl.Workbook(); ws = wb.active; ws.title = f'Cycle {CYCLE} regression'
bold = Font(bold=True)
ws.append([f'Harvest cycle {CYCLE}: lab value = slope × SPAD + intercept'])
ws['A1'].font = bold
ws.append(['SPAD source', 'Lab measure', 'Samples', 'n', 'Slope', 'Intercept', 'Pearson r', 'R²', 'p-value'])
for c in ws[2]: c.font = bold; c.alignment = Alignment(horizontal='center', wrap_text=True)
for r in rec:
    ws.append(r[:4] + [round(r[4], 4), round(r[5], 3), round(r[6], 3), round(r[7], 3), float(f'{r[8]:.3g}')])
for j, w in enumerate((13, 18, 24, 6, 10, 10, 10, 9, 11), 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = 'A3'; ws.auto_filter.ref = f'A2:I{ws.max_row}'
wb.save(TABLES / 'SPAD_Cycle5_Regression.xlsx')
for r in rec: print(f'{r[0]:11s} {r[1]:18s} {r[2]:24s} n={r[3]:2d} y={r[4]:.3f}x{r[5]:+.2f} r={r[6]:.2f} R2={r[7]:.2f} p={r[8]:.2g}')

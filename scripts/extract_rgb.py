"""Extract the mean leaf colour (R, G, B, 0-255) of every preprocessed image into Excel.

Reads  data_preprocessed/cycle-N/<sample>.png (transparent background)
Writes Leaf_RGB_Data.xlsx, one sheet per harvest cycle, rows in the same order as SPAD_Data.xlsx.
Only leaf pixels count: fully opaque pixels, minus the orange rubber-band lines that cross the leaves.
"""
from pathlib import Path

import cv2
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent          # project folder
DATA = ROOT / "data"                                        # measurement tables
TABLES = ROOT / "results" / "tables"
FIGS = ROOT / "results" / "figures"
SRC = ROOT / "data_preprocessed"


def leaf_rgb(path):
    bgra = cv2.imread(str(path), cv2.IMREAD_UNCHANGED)
    hsv = cv2.cvtColor(bgra[..., :3], cv2.COLOR_BGR2HSV)
    leaf = bgra[..., 3] == 255
    leaf = cv2.erode(leaf.astype(np.uint8), np.ones((5, 5), np.uint8)) > 0     # drop the blended outline
    band = cv2.dilate(cv2.inRange(hsv, (3, 90, 90), (22, 255, 255)), np.ones((3, 3), np.uint8)) > 0
    leaf &= ~band
    r, g, b = (bgra[..., c][leaf].astype(float) for c in (2, 1, 0))            # OpenCV stores B, G, R
    return r.mean(), g.mean(), b.mean(), int(leaf.sum())


def main():
    spad = openpyxl.load_workbook(DATA / "SPAD_Data.xlsx", data_only=True)
    out = openpyxl.Workbook()
    out.remove(out.active)
    for cycle, sheet in enumerate(spad.worksheets, 1):
        ws = out.create_sheet(f"c_{cycle}")
        ws.append(["", "Sample", "R", "G", "B", "Leaf pixels"])
        for r in range(3, sheet.max_row + 1):
            idx, label = sheet.cell(r, 1).value, sheet.cell(r, 2).value
            f = SRC / f"cycle-{cycle}" / f"{label}.png"
            if f.exists():
                R, G, B, n = leaf_rgb(f)
                ws.append([idx, label, round(R, 3), round(G, 3), round(B, 3), n])
            else:
                ws.append([idx, label, None, None, None, None])                # no image for this sample
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
            for c in row:
                c.number_format = "0.00"
        for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
            for c in row:
                c.number_format = "#,##0"
        for col, w in zip(range(1, 7), (6, 14, 10, 10, 10, 13)):
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = "C2"
    out.save(TABLES / "Leaf_RGB_Data.xlsx")


if __name__ == "__main__":
    main()

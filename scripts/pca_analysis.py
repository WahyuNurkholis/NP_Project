"""Does PCA help? What the main components capture, and whether PCA before regression improves cross-cycle prediction.

Outputs results/tables/PCA_analysis.xlsx and results/figures/pca_scores.svg (spectra and camera-colour PCA score plots).
PCA is fitted on the training cycles only inside every validation fold, so nothing leaks.
"""
import os
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import dnn_feasibility as dnf               # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
TABLES, FIGS = ROOT / "results" / "tables", ROOT / "results" / "figures"


def pca_fit(X, k=None):
    mu, sd = X.mean(0), X.std(0) + 1e-9
    Z = (X - mu) / sd
    U, S, Vt = np.linalg.svd(Z, full_matrices=False)
    var = S ** 2 / (S ** 2).sum()
    return mu, sd, Vt[: (k or len(S))], var


def eta2(scores, labels):
    """Share of each component's variance explained by a grouping (one-way ANOVA eta squared)."""
    out = []
    for j in range(scores.shape[1]):
        s = scores[:, j]; tot = ((s - s.mean()) ** 2).sum()
        between = sum(len(s[labels == v]) * (s[labels == v].mean() - s.mean()) ** 2 for v in np.unique(labels))
        out.append(between / tot)
    return np.array(out)


def loco_pca(X, y, g, task, kind, k):
    pred = np.empty(len(y), dtype=float)
    for h in np.unique(g):
        te, tr = g == h, g != h
        if k is None:
            Xtr, Xte = X[tr], X[te]
        else:
            mu, sd, V, _ = pca_fit(X[tr], k)
            Xtr, Xte = ((X[tr] - mu) / sd) @ V.T, ((X[te] - mu) / sd) @ V.T
        pred[te] = dnf.fit_predict(Xtr, y[tr], Xte, task, kind)
    return (1 - ((y - pred) ** 2).sum() / ((y - y.mean()) ** 2).sum()) if task == "reg" else (pred == y).mean()


def main():
    df = dnf.build()
    sets = {"Spectra 400-850 nm": (df[df.QC_flag == "ok"].reset_index(drop=True), "spec"),
            "Camera colour": (df[[np.isfinite(v).all() for v in df["rgb"]]].reset_index(drop=True), "rgb")}
    struct, cvrows, scores_all = [], [], {}
    for name, (d, col) in sets.items():
        X = np.vstack(d[col].values); g = d.Harvest.values
        mu, sd, V, var = pca_fit(X, 10)
        S = ((X - mu) / sd) @ V.T; scores_all[name] = (S, var, d)
        for j in range(6):
            struct.append(dict(features=name, component=f"PC{j + 1}", variance_pct=100 * var[j], cumulative_pct=100 * var[: j + 1].sum(),
                               share_harvest_cycle=eta2(S[:, [j]], g)[0], share_shade_block=eta2(S[:, [j]], d.Block.values)[0], share_irrigation=eta2(S[:, [j]], d["T"].values)[0]))
        for tgt, task in (("SPAD", "reg"), ("CHL_A", "reg"), ("CHL_B", "reg"), ("T", "clf"), ("Block", "clf")):
            y = d[tgt].values.astype(float)
            for kind in ("ridge", "mlp"):
                row = dict(features=name, target={"T": "Irrigation level", "Block": "Shade block"}.get(tgt, tgt), metric="R²" if task == "reg" else "accuracy", model=kind)
                for k in (None, 3, 5, 10):
                    row["no PCA" if k is None else f"PCA {k} comp."] = loco_pca(X, y, g, task, kind, k)
                cvrows.append(row); print(name, row["target"], kind, {k: round(v, 2) for k, v in row.items() if k.startswith(("no", "PCA"))}, flush=True)
    st, cv = pd.DataFrame(struct), pd.DataFrame(cvrows)
    wb = Workbook(); wb.remove(wb.active); bold = Font(bold=True)
    for title, t, note in (("What the components capture", st, "Variance explained by each PCA component, and how much of it is explained by harvest cycle, shade block and irrigation level (share of the component's variance)."),
                           ("Cross-cycle prediction", cv, "Leave-one-harvest-cycle-out. PCA fitted on the training cycles only. R² < 0 = worse than predicting the mean; accuracy for the classes (chance: irrigation 0.33-0.36, block 0.5).")):
        ws = wb.create_sheet(title); ws.append([note]); ws["A1"].font = bold; ws.append(list(t.columns))
        for c in ws[2]: c.font = bold; c.alignment = Alignment(horizontal="center", wrap_text=True)
        for _, r in t.iterrows(): ws.append([round(float(v), 3) if isinstance(v, (float, np.floating)) else v for v in r.values])
        for i in range(1, len(t.columns) + 1): ws.column_dimensions[get_column_letter(i)].width = 20
    wb.save(TABLES / "PCA_analysis.xlsx")
    print(st.round(2).to_string(index=False))

    # ---- figure: PCA score plots
    SURF, INK, INK2, GRID = "#fcfcfb", "#0b0b0b", "#3d3c39", "#e2e1dc"
    plt.rcParams["svg.fonttype"] = "path"
    plt.rcParams.update({"font.family": "serif", "font.serif": ["Times New Roman", "Nimbus Roman", "STIXGeneral"], "text.color": INK,
                         "axes.labelcolor": INK, "xtick.color": INK2, "ytick.color": INK2, "font.size": 15})
    CYC = {1: ("#2a78d6", "o"), 2: ("#eb6834", "s"), 3: ("#1baf7a", "^"), 4: ("#eda100", "D"), 5: ("#e87ba4", "v")}
    BLK = {1: ("#2a78d6", "o", "B1 · outside shade"), 2: ("#eb6834", "s", "B2 · under panel")}
    TRT = {1: ("#2a78d6", "o", "T1 · 100%"), 2: ("#eb6834", "s", "T2 · 75%"), 3: ("#1baf7a", "^", "T3 · 50%")}
    fig, axes = plt.subplots(2, 3, figsize=(16, 10.5), facecolor=SURF)
    fig.subplots_adjust(left=0.07, right=0.985, top=0.86, bottom=0.075, hspace=0.42, wspace=0.26)
    fig.text(0.07, 0.972, "PCA of the spectra and camera colour: what separates the samples", fontsize=25, fontweight="bold", va="top")
    fig.text(0.07, 0.925, "Each point is one sample on the first two principal components. If the groups do not separate, the main variation in the data is not about that factor.",
             fontsize=15, color=INK2, va="top")
    for i, (name, (S, var, d)) in enumerate(scores_all.items()):
        for j, (fac, key, spec) in enumerate((("cycle", "Harvest", None), ("shade", "Block", BLK), ("irrigation", "T", TRT))):
            ax = axes[i, j]; ax.set_facecolor(SURF)
            vals = d[key].values
            for v in np.unique(vals):
                col, mk = (CYC[v] if spec is None else spec[v][:2]); m = vals == v
                ax.scatter(S[m, 0], S[m, 1], s=44, marker=mk, color=col, alpha=0.75, linewidths=0, zorder=3)
            ax.set_title(f"{name.replace(' 400-850 nm', '')} · by {fac}", loc="left", fontsize=17, fontweight="bold", pad=8)
            ax.set_xlabel(f"PC1 ({100 * var[0]:.0f}% of variance)", fontsize=15); ax.set_ylabel(f"PC2 ({100 * var[1]:.0f}%)", fontsize=15)
            ax.grid(color=GRID, lw=1); ax.set_axisbelow(True)
            for s in ("top", "right"): ax.spines[s].set_visible(False)
            for s in ("left", "bottom"): ax.spines[s].set_color("#b9b8b2")
            ax.tick_params(length=0, labelsize=13.5)
            lo, hi = ax.get_ylim(); ax.set_ylim(lo, hi + (hi - lo) * 0.3)
            h = [Line2D([], [], color=CYC[v][0], marker=CYC[v][1], ls="", ms=8, label=f"Cycle {v}") for v in CYC] if spec is None else \
                [Line2D([], [], color=spec[v][0], marker=spec[v][1], ls="", ms=8, label=spec[v][2]) for v in spec]
            ax.legend(handles=h, loc="upper left", frameon=False, fontsize=12, ncol=3, columnspacing=1.0, handletextpad=0.3)
    fig.savefig(FIGS / "pca_scores.svg", facecolor=SURF)
    if os.environ.get("PREVIEW_DIR"):
        fig.savefig(Path(os.environ["PREVIEW_DIR"]) / "pca_scores.png", dpi=int(os.environ.get('PREVIEW_DPI', 100)), facecolor=SURF)


if __name__ == "__main__":
    main()

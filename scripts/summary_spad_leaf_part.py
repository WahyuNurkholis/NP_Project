"""Summary of the linear fit of lab chlorophyll on SPAD, for each SPAD source:
whole leaf (SPAD_Data.xlsx) and the up / mid / low parts (backup_original/SPAD_Data_B3_dropped_per_part.xlsx).

Writes results/tables/SPAD_Leaf_Part_Summary.xlsx:
  R2 summary   R² per subset (rows) and SPAD source × lab measure (columns), best source per row in bold
  Full stats   n, slope, intercept, r, R², p-value for every source, subset and lab measure
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy import stats

ROOT = Path(__file__).resolve().parent.parent          # project folder
DATA = ROOT / 'data'
TABLES = ROOT / 'results' / 'tables'

# ---- SPAD per (cycle, sample) for each source; means over the readings present
SPAD = {'Whole leaf': {}}
for cyc, ws in enumerate(openpyxl.load_workbook(DATA / 'SPAD_Data.xlsx', data_only=True).worksheets, 1):
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 2).value is not None: SPAD['Whole leaf'][(cyc, ws.cell(r, 2).value)] = ws.cell(r, 3).value
parts = openpyxl.load_workbook(DATA / 'backup_original' / 'SPAD_Data_B3_dropped_per_part.xlsx', data_only=True).worksheets
for part, label in (('up', 'Up'), ('mid', 'Mid'), ('low', 'Low')):
    SPAD[label] = {}
    for cyc, ws in enumerate(parts, 1):
        cols = [j for j in range(3, ws.max_column + 1) if ws.cell(1, j).value == part]
        for r in range(3, ws.max_row + 1):
            v = [ws.cell(r, j).value for j in cols]; v = [x for x in v if isinstance(x, (int, float))]
            if ws.cell(r, 2).value is not None and v: SPAD[label][(cyc, ws.cell(r, 2).value)] = sum(v) / len(v)

lab = [(int(r[0]), r[1], r[2], r[3], r[4]) for r in openpyxl.load_workbook(DATA / 'Chl_Lab_Data.xlsx', data_only=True).active
       .iter_rows(min_row=2, max_col=5, values_only=True) if r[1] is not None]
MEAS = (('Chl a', 2), ('Chl b', 3), ('Total Chl', 4))
SUBSETS = [('All samples', lambda h, s: True), ('B1 outside panel shade', lambda h, s: 'B1' in s),
           ('B2 under solar panel', lambda h, s: 'B2' in s)] + [(f'Cycle {c}', lambda h, s, c=c: h == c) for c in range(1, 6)]

res = {}                                   # (source, subset, measure) -> stats
for src, sp in SPAD.items():
    for sub, f in SUBSETS:
        rows = [r for r in lab if f(r[0], r[1])]
        for m, i in MEAS:
            lr = stats.linregress([sp[(r[0], r[1])] for r in rows], [r[i] for r in rows])
            res[(src, sub, m)] = dict(n=len(rows), slope=lr.slope, intercept=lr.intercept, r=lr.rvalue, R2=lr.rvalue ** 2, p=lr.pvalue)

# ---- workbook
bold, thin = Font(bold=True), Side(style='thin', color='B9B8B2')
ctr = Alignment(horizontal='center', vertical='center', wrap_text=True)
best_fill = PatternFill('solid', fgColor='CDE2FB')
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'R2 summary'
ws.append(['R² of the linear fit: lab chlorophyll = slope × SPAD + intercept. Best SPAD source per row and lab measure is bold and shaded.'])
ws['A1'].font = bold
ws.append(['Subset', 'n'] + [m for m, _ in MEAS for _ in SPAD])
ws.append(['', ''] + [src for _ in MEAS for src in SPAD])
for k, (m, _) in enumerate(MEAS):
    c0 = 3 + k * len(SPAD); ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + len(SPAD) - 1)
ws.merge_cells('A2:A3'); ws.merge_cells('B2:B3')
for row in ws.iter_rows(min_row=2, max_row=3):
    for c in row: c.font = bold; c.alignment = ctr; c.border = Border(bottom=thin)
for sub, _ in SUBSETS:
    vals = [res[(src, sub, m)]['R2'] for m, _ in MEAS for src in SPAD]
    ws.append([sub, res[('Whole leaf', sub, 'Chl a')]['n']] + [round(v, 3) for v in vals])
    rr = ws.max_row
    for k in range(len(MEAS)):
        block = vals[k * len(SPAD):(k + 1) * len(SPAD)]; j = 3 + k * len(SPAD) + block.index(max(block))
        ws.cell(rr, j).font = bold; ws.cell(rr, j).fill = best_fill
    for c in ws[rr][1:]: c.alignment = ctr
ws.append([])
ws.append(['Whole leaf = SPAD_Data.xlsx (whole-leaf mean). Up / Mid / Low = mean of that leaf part in SPAD_Data_B3_dropped_per_part.xlsx, over the readings present.'])
ws.append(['Lab values from Chl_Lab_Data.xlsx; 270 samples (54 samples × 5 harvest cycles) matched for every source.'])
ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 7
for j in range(3, 3 + len(MEAS) * len(SPAD)): ws.column_dimensions[get_column_letter(j)].width = 11
ws.freeze_panes = 'C4'

ws = wb.create_sheet('Full stats')
ws.append(['SPAD source', 'Subset', 'Lab measure', 'n', 'Slope', 'Intercept', 'Pearson r', 'R²', 'p-value'])
for c in ws[1]: c.font = bold; c.alignment = ctr
for src in SPAD:
    for sub, _ in SUBSETS:
        for m, _ in MEAS:
            s = res[(src, sub, m)]
            ws.append([src, sub, m, s['n'], round(s['slope'], 4), round(s['intercept'], 3), round(s['r'], 3), round(s['R2'], 3), float(f"{s['p']:.3g}")])
for j, w in enumerate((12, 24, 12, 7, 10, 10, 10, 9, 11), 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions

out = TABLES / 'SPAD_Leaf_Part_Summary.xlsx'; wb.save(out); print('wrote', out)
for m, _ in MEAS:
    print('\n' + m); print('Subset'.ljust(24) + ''.join(s.rjust(12) for s in SPAD))
    for sub, _ in SUBSETS: print(sub.ljust(24) + ''.join(f"{res[(s, sub, m)]['R2']:12.2f}" for s in SPAD))

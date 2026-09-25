"""Lower-leaf SPAD (mean of the 'low' columns in SPAD_Data_B3_dropped_per_part.xlsx) linked to the lab chlorophyll data.

Writes results/tables/SPAD_Low_vs_Chl_Lab.xlsx with two sheets:
  Data        one row per sample and harvest cycle: low SPAD readings, their mean, and the lab values
  Regression  linear fit of each lab measure on low SPAD, for all samples, each block and each cycle
"""
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from scipy import stats

ROOT = Path(__file__).resolve().parent.parent          # project folder
DATA = ROOT / 'data'
TABLES = ROOT / 'results' / 'tables'

# low SPAD readings, kept per reading; mean over the readings present
low = {}
for cyc, ws in enumerate(openpyxl.load_workbook(DATA / 'backup_original' / 'SPAD_Data_B3_dropped_per_part.xlsx', data_only=True).worksheets, 1):
    cols = [j for j in range(3, ws.max_column + 1) if ws.cell(1, j).value == 'low']
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 2).value is not None:
            low[(cyc, ws.cell(r, 2).value)] = [ws.cell(r, j).value if isinstance(ws.cell(r, j).value, (int, float)) else None for j in cols]

rows = []
for r in openpyxl.load_workbook(DATA / 'Chl_Lab_Data.xlsx', data_only=True).active.iter_rows(min_row=2, max_col=5, values_only=True):
    if r[1] is None:
        continue
    t, b = re.match(r'T(\d)B(\d)', r[1]).groups()
    v = low[(int(r[0]), r[1])]; p = [x for x in v if x is not None]
    rows.append(dict(Harvest=int(r[0]), Sample=r[1], T=int(t), Block=int(b), low=v, n=len(p), SPAD=sum(p) / len(p),
                     A=r[2], B=r[3], Tot=r[4]))

wb = openpyxl.Workbook(); bold = Font(bold=True); ctr = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws = wb.active; ws.title = 'Data'
nmax = max(len(d['low']) for d in rows)
hdr = ['Harvest', 'Sample', 'Treatment', 'Block'] + [f'Low SPAD {i}' for i in range(1, nmax + 1)] + \
      ['Readings used', 'Low SPAD mean', 'Lab Chl a', 'Lab Chl b', 'Lab Total Chl']
ws.append(hdr)
for d in rows:
    ws.append([d['Harvest'], d['Sample'], f"T{d['T']}", f"B{d['Block']}"] + d['low'] + [None] * (nmax - len(d['low'])) +
              [d['n'], round(d['SPAD'], 3), d['A'], d['B'], d['Tot']])
for c in ws[1]: c.font = bold; c.alignment = ctr
ws.freeze_panes = 'C2'
for i in range(1, len(hdr) + 1): ws.column_dimensions[get_column_letter(i)].width = 12
ws.column_dimensions['B'].width = 14

ws = wb.create_sheet('Regression')
ws.append(['Linear regression: lab value = slope × low SPAD + intercept'])
ws['A1'].font = bold
ws.append(['Subset', 'Lab measure', 'n', 'Slope', 'Intercept', 'Pearson r', 'R²', 'p-value'])
for c in ws[2]: c.font = bold; c.alignment = ctr
subsets = [('All samples', lambda d: True), ('B1 outside panel shade', lambda d: d['Block'] == 1),
           ('B2 under solar panel', lambda d: d['Block'] == 2)] + \
          [(f'Cycle {h}', lambda d, h=h: d['Harvest'] == h) for h in range(1, 6)]
for name, f in subsets:
    s = [d for d in rows if f(d)]
    for key, lab in (('A', 'Chl a'), ('B', 'Chl b'), ('Tot', 'Total Chl')):
        lr = stats.linregress([d['SPAD'] for d in s], [d[key] for d in s])
        ws.append([name, lab, len(s), round(lr.slope, 4), round(lr.intercept, 3), round(lr.rvalue, 3), round(lr.rvalue ** 2, 3), float(f'{lr.pvalue:.3g}')])
ws.column_dimensions['A'].width = 24
for i in range(2, 9): ws.column_dimensions[get_column_letter(i)].width = 13

out = TABLES / 'SPAD_Low_vs_Chl_Lab.xlsx'
wb.save(out)
print('wrote', out, len(rows), 'rows')

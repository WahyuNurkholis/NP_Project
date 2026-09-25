"""Which analyses can show a real, defensible result with these data?

A. Irrigation effect as a paired difference to the well-watered plot (T - T1, same block, replication, repeat and cycle).
B. Camera colour calibrated inside each harvest cycle, validated by leaving out one replication at a time.
C. Shade detection (B1 vs B2) from SPAD or camera colour, validated by leaving out one harvest cycle.
Writes results/tables/Promising_analyses.xlsx.
"""
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from scipy import stats

sys.path.insert(0, str(Path(__file__).resolve().parent))
import spectrometer_analysis as sa          # noqa: E402

DATA, TABLES = sa.DATA, sa.TABLES


def load():
    spad = {}
    import openpyxl
    for c, ws in enumerate(openpyxl.load_workbook(DATA / "SPAD_Data.xlsx", data_only=True).worksheets, 1):
        for r in range(3, ws.max_row + 1):
            spad[(c, ws.cell(r, 2).value)] = ws.cell(r, 3).value
    rows = []
    for r in openpyxl.load_workbook(DATA / "Chl_Lab_Data.xlsx", data_only=True).active.iter_rows(min_row=2, max_col=5, values_only=True):
        if r[1] is None: continue
        t, b, rep, x = re.match(r"T(\d)B(\d)R(\d)\((\d)\)", r[1]).groups()
        rows.append(dict(Harvest=int(r[0]), Sample=r[1], T=int(t), Block=int(b), Rep=int(rep), X=int(x),
                         SPAD=spad[(int(r[0]), r[1])], CHL_A=r[2], CHL_B=r[3]))
    df = pd.DataFrame(rows)
    vi = pd.read_excel(TABLES / "Vegetation_Indices.xlsx", sheet_name="Indices")
    return df.merge(vi[["Harvest", "Sample", "DGCI", "Hue (partial)", "NDVI green", "VARI", "r", "g", "b"]], on=["Harvest", "Sample"], how="left")


def paired_effect(df, y):
    rec = []
    for (h, b), d in df.groupby(["Harvest", "Block"]):
        p = d.pivot_table(index=["Rep", "X"], columns="T", values=y)
        for t in (2, 3):
            diff = (p[t] - p[1]).dropna()
            tt = stats.ttest_1samp(diff, 0)
            rec.append(dict(target=y, harvest=h, block=b, treatment=f"T{t} - T1", n=len(diff), mean_diff=diff.mean(), rel_pct=100 * diff.mean() / p[1].mean(), p=tt.pvalue))
    return pd.DataFrame(rec)


def within_cycle_calibration(df, feats, y):
    """Fit y ~ features inside one cycle, leaving out one replication at a time."""
    rec = []
    for h, d in df.groupby("Harvest"):
        d = d.dropna(subset=feats + [y]); pred = np.zeros(len(d))
        X = d[feats].values; t = d[y].values
        for rep in (1, 2, 3):
            te = (d.Rep == rep).values; tr = ~te
            mu, sd = X[tr].mean(0), X[tr].std(0) + 1e-9
            A = np.c_[np.ones(tr.sum()), (X[tr] - mu) / sd]
            beta = np.linalg.solve(A.T @ A + 5 * np.eye(A.shape[1]), A.T @ t[tr])
            pred[te] = np.c_[np.ones(te.sum()), (X[te] - mu) / sd] @ beta
        rec.append(dict(target=y, harvest=h, n=len(d), R2=1 - ((t - pred) ** 2).sum() / ((t - t.mean()) ** 2).sum(), r=stats.pearsonr(pred, t).statistic))
    return pd.DataFrame(rec)


def shade_detection(df, feats, name):
    d = df.dropna(subset=feats); X = d[feats].values; y = (d.Block == 2).astype(float).values; g = d.Harvest.values; acc = []
    for h in np.unique(g):
        te, tr = g == h, g != h
        mu, sd = X[tr].mean(0), X[tr].std(0) + 1e-9
        A = np.c_[np.ones(tr.sum()), (X[tr] - mu) / sd]
        beta = np.linalg.solve(A.T @ A + 5 * np.eye(A.shape[1]), A.T @ (2 * y[tr] - 1))
        acc.append((((np.c_[np.ones(te.sum()), (X[te] - mu) / sd] @ beta) > 0) == (y[te] == 1)).mean())
    return dict(features=name, accuracy_mean=np.mean(acc), accuracy_min=np.min(acc), accuracy_max=np.max(acc), chance=0.5)


def main():
    df = load()
    A = pd.concat([paired_effect(df, y) for y in ("SPAD", "CHL_A", "CHL_B")])
    B = pd.concat([within_cycle_calibration(df, ["DGCI", "r", "Hue (partial)", "NDVI green"], y) for y in ("SPAD", "CHL_A")])
    C = pd.DataFrame([shade_detection(df, ["SPAD"], "SPAD"), shade_detection(df, ["DGCI", "r", "Hue (partial)", "NDVI green", "VARI"], "Camera colour indices"),
                      shade_detection(df, ["DGCI", "r", "Hue (partial)", "NDVI green", "VARI", "SPAD"], "Camera + SPAD")])
    wb = Workbook(); wb.remove(wb.active); bold = Font(bold=True)
    for name, t, note in (("A paired irrigation effect", A, "Paired difference to the well-watered plot T1 (same block, replication, repeat, cycle). rel_pct = % of the T1 mean."),
                          ("B calibration within cycle", B, "Camera colour indices -> SPAD / chl a, fitted inside one harvest cycle, one replication left out at a time."),
                          ("C shade detection", C, "Classify B1 vs B2, trained on 4 cycles and tested on the 5th.")):
        ws = wb.create_sheet(name); ws.append([note]); ws["A1"].font = bold; ws.append(list(t.columns))
        for c in ws[2]: c.font = bold; c.alignment = Alignment(horizontal="center")
        for _, r in t.iterrows(): ws.append([round(float(v), 4) if isinstance(v, (float, np.floating)) else (int(v) if isinstance(v, (int, np.integer)) else v) for v in r.values])
        for i in range(1, len(t.columns) + 1): ws.column_dimensions[get_column_letter(i)].width = 18
    wb.save(TABLES / "Promising_analyses.xlsx")
    pd.set_option("display.width", 200)
    a = A[A.target == "SPAD"]
    print("A. SPAD, T - T1 paired:\n", a.round(3).to_string(index=False))
    print("\nA. share of block x cycle cells where T3<T1:", {y: f"{(A[(A.target==y)&(A.treatment=='T3 - T1')].mean_diff<0).sum()}/10" for y in ('SPAD','CHL_A','CHL_B')})
    print("   significant (p<0.05) cells T3-T1:", {y: int((A[(A.target==y)&(A.treatment=='T3 - T1')].p<0.05).sum()) for y in ('SPAD','CHL_A','CHL_B')}, "of 10")
    print("\nB. within-cycle calibration:\n", B.round(2).to_string(index=False))
    print("\nC. shade detection:\n", C.round(2).to_string(index=False))


if __name__ == "__main__":
    main()

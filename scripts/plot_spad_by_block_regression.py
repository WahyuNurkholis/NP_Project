"""SPAD vs lab chlorophyll with the two shade blocks analysed separately
(B1 outside panel shade, B2 under solar panel), for each SPAD source
(whole leaf from SPAD_Data.xlsx; up / mid / low from backup_original/SPAD_Data_B3_dropped_per_part.xlsx).

Writes
  results/figures/spad_cycle5_B1_regression.svg, results/figures/spad_cycle5_B2_regression.svg
      cycle 5 only, one figure per block (rows = lab measure, columns = SPAD source); PNGs to $PREVIEW_DIR if set
  results/tables/SPAD_By_Block_Regression.xlsx
      linear fit per block × SPAD source × lab measure, for each cycle and for all cycles pooled
"""
import os
import re
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.lines import Line2D
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy import stats

ROOT = Path(__file__).resolve().parent.parent          # project folder
DATA = ROOT / 'data'
FIGS = ROOT / 'results' / 'figures'
TABLES = ROOT / 'results' / 'tables'
FIG_CYCLE = 5

# ---- SPAD per (cycle, sample) for each source; part means over the readings present
SPAD = {'Whole leaf': {}}
for cyc, ws in enumerate(openpyxl.load_workbook(DATA / 'SPAD_Data.xlsx', data_only=True).worksheets, 1):
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 2).value is not None: SPAD['Whole leaf'][(cyc, ws.cell(r, 2).value)] = ws.cell(r, 3).value
parts = openpyxl.load_workbook(DATA / 'backup_original' / 'SPAD_Data_B3_dropped_per_part.xlsx', data_only=True).worksheets
for part, label in (('up', 'Upper leaf'), ('mid', 'Mid leaf'), ('low', 'Lower leaf')):
    SPAD[label] = {}
    for cyc, ws in enumerate(parts, 1):
        cols = [j for j in range(3, ws.max_column + 1) if ws.cell(1, j).value == part]
        for r in range(3, ws.max_row + 1):
            v = [ws.cell(r, j).value for j in cols]; v = [x for x in v if isinstance(x, (int, float))]
            if ws.cell(r, 2).value is not None and v: SPAD[label][(cyc, ws.cell(r, 2).value)] = sum(v) / len(v)

LAB = [dict(cyc=int(r[0]), s=r[1], blk=int(re.match(r'T\dB(\d)', r[1]).group(1)), a=r[2], b=r[3], t=r[4])
       for r in openpyxl.load_workbook(DATA / 'Chl_Lab_Data.xlsx', data_only=True).active.iter_rows(min_row=2, max_col=5, values_only=True)
       if r[1] is not None]
MEAS = (('Chlorophyll a', 'a'), ('Chlorophyll b', 'b'), ('Total chlorophyll', 't'))
BNAME = {1: 'B1 – outside panel shade', 2: 'B2 – under solar panel'}


def fit(rows, src, key):
    x = np.array([SPAD[src][(d['cyc'], d['s'])] for d in rows], float); y = np.array([d[key] for d in rows], float)
    return x, y, stats.linregress(x, y)


# ---- figures: cycle 5, one per block
SURF, INK, INK2, GRID = '#fcfcfb', '#0b0b0b', '#3d3c39', '#e2e1dc'
BC = {1: '#2a78d6', 2: '#eb6834'}; BM = {1: 'o', 2: 's'}
plt.rcParams['svg.fonttype'] = 'path'
plt.rcParams.update({'font.family': 'serif', 'font.serif': ['Times New Roman', 'Nimbus Roman', 'STIXGeneral'],
                     'mathtext.fontset': 'stix', 'text.color': INK, 'axes.labelcolor': INK,
                     'xtick.color': INK2, 'ytick.color': INK2, 'font.size': 15})
for b in (1, 2):
    rows = [d for d in LAB if d['cyc'] == FIG_CYCLE and d['blk'] == b]
    fig, axes = plt.subplots(3, 4, figsize=(22, 17), facecolor=SURF)
    fig.subplots_adjust(left=0.06, right=0.99, top=0.885, bottom=0.09, hspace=0.42, wspace=0.28)
    fig.text(0.06, 0.975, f'SPAD vs laboratory chlorophyll, cycle {FIG_CYCLE}, {BNAME[b]} only', fontsize=28, fontweight='bold', va='top')
    fig.text(0.06, 0.94, f'Each point is one sample (n = {len(rows)}: 3 treatments × 3 replicates × 3 subsamples). '
             'Line = linear fit within this block. Columns: SPAD source; rows: laboratory measure.', fontsize=16, color=INK2, va='top')
    for i, (mname, key) in enumerate(MEAS):
        for j, src in enumerate(SPAD):
            ax = axes[i, j]; ax.set_facecolor(SURF)
            x, y, lr = fit(rows, src, key)
            ax.scatter(x, y, s=64, marker=BM[b], color=BC[b], alpha=0.8, linewidths=0, zorder=3)
            xs = np.array([x.min(), x.max()])
            ax.plot(xs, lr.intercept + lr.slope * xs, color=INK, lw=2.4, zorder=4)
            sign = '+' if lr.intercept >= 0 else '−'
            ptxt = f'p = {lr.pvalue:.3f}' if lr.pvalue >= 0.001 else 'p < 0.001'
            ax.legend(handles=[Line2D([], [], color=BC[b], marker=BM[b], ls='', ms=9, label=BNAME[b]),
                               Line2D([], [], color=INK, lw=2.4, label=f'Fit: y = {lr.slope:.2f}x {sign} {abs(lr.intercept):.1f}\n'
                                                                         f'r = {lr.rvalue:.2f}, R² = {lr.rvalue**2:.2f}, {ptxt}')],
                      loc='upper left', frameon=True, facecolor=SURF, edgecolor='none', framealpha=0.9,
                      fontsize=12.5, handletextpad=0.5, borderpad=0.3, labelspacing=0.35)
            lo, hi = y.min(), y.max(); ax.set_ylim(max(0, lo - 0.1 * (hi - lo)), hi + 0.75 * (hi - lo))
            ax.set_title(f'{mname} – {src} SPAD', loc='left', fontsize=17, fontweight='bold', pad=8)
            ax.set_xlabel(f'{src} SPAD value', fontsize=15, labelpad=5)
            ax.set_ylabel(f'{mname} (lab value)', fontsize=15, labelpad=5)
            ax.grid(color=GRID, lw=1); ax.set_axisbelow(True)
            for s in ('top', 'right'): ax.spines[s].set_visible(False)
            for s in ('left', 'bottom'): ax.spines[s].set_color('#b9b8b2')
            ax.tick_params(length=0, labelsize=13.5)
    fig.text(0.06, 0.022, 'Whole leaf = SPAD_Data.xlsx; upper / mid / lower leaf = mean of that part in SPAD_Data_B3_dropped_per_part.xlsx '
             '(readings present only). Axes are scaled to this block. Font: Nimbus Roman (Times-compatible).', fontsize=14, color=INK2, style='italic')
    name = f'spad_cycle{FIG_CYCLE}_B{b}_regression'
    fig.savefig(FIGS / f'{name}.svg', facecolor=SURF)
    if os.environ.get('PREVIEW_DIR'):
        fig.savefig(Path(os.environ['PREVIEW_DIR']) / f'{name}.png', dpi=int(os.environ.get('PREVIEW_DPI', 100)), facecolor=SURF)
    plt.close(fig)

# ---- table: every block × cycle (plus pooled) × source × measure
rec = []
for b in (1, 2):
    for cyc in ('All cycles', 1, 2, 3, 4, 5):
        rows = [d for d in LAB if d['blk'] == b and (cyc == 'All cycles' or d['cyc'] == cyc)]
        for mname, key in MEAS:
            for src in SPAD:
                _, _, lr = fit(rows, src, key)
                rec.append([BNAME[b], cyc if cyc == 'All cycles' else f'Cycle {cyc}', mname, src, len(rows),
                            lr.slope, lr.intercept, lr.rvalue, lr.rvalue ** 2, lr.pvalue])
bold = Font(bold=True); sig = PatternFill('solid', fgColor='CDE2FB')
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'By block'
ws.append(['Lab value = slope × SPAD + intercept, fitted separately in each block. Rows with p < 0.05 are shaded.'])
ws['A1'].font = bold
ws.append(['Block', 'Cycle', 'Lab measure', 'SPAD source', 'n', 'Slope', 'Intercept', 'Pearson r', 'R²', 'p-value'])
for c in ws[2]: c.font = bold; c.alignment = Alignment(horizontal='center', wrap_text=True)
for r in rec:
    ws.append(r[:5] + [round(r[5], 4), round(r[6], 3), round(r[7], 3), round(r[8], 3), float(f'{r[9]:.3g}')])
    if r[9] < 0.05:
        for c in ws[ws.max_row]: c.fill = sig
for j, w in enumerate((26, 11, 18, 12, 6, 10, 10, 10, 9, 11), 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = 'A3'; ws.auto_filter.ref = f'A2:J{ws.max_row}'

ws = wb.create_sheet('R2 overview')
ws.append(['R² per block and cycle (rows) × lab measure and SPAD source (columns). Bold = p < 0.05.']); ws['A1'].font = bold
ws.append(['Block', 'Cycle', 'n'] + [f'{m}\n{s}' for m, _ in MEAS for s in SPAD])
for c in ws[2]: c.font = bold; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[2].height = 45
idx = {(r[0], r[1], r[2], r[3]): r for r in rec}
for b in (1, 2):
    for cyc in ('All cycles', 'Cycle 1', 'Cycle 2', 'Cycle 3', 'Cycle 4', 'Cycle 5'):
        cells = [idx[(BNAME[b], cyc, m, s)] for m, _ in MEAS for s in SPAD]
        ws.append([BNAME[b], cyc, cells[0][4]] + [round(r[8], 3) for r in cells])
        for k, r in enumerate(cells):
            if r[9] < 0.05: ws.cell(ws.max_row, 4 + k).font = bold
ws.column_dimensions['A'].width = 26; ws.column_dimensions['B'].width = 11; ws.column_dimensions['C'].width = 6
for j in range(4, 4 + 12): ws.column_dimensions[get_column_letter(j)].width = 12
ws.freeze_panes = 'D3'
wb.save(TABLES / 'SPAD_By_Block_Regression.xlsx')

for b in (1, 2):
    print('\n' + BNAME[b])
    for cyc in ('All cycles', 'Cycle 5'):
        for m, _ in MEAS:
            print(f'  {cyc:10s} {m:18s}', '  '.join(f"{s.split()[0]:5s} R2={idx[(BNAME[b], cyc, m, s)][8]:.2f} p={idx[(BNAME[b], cyc, m, s)][9]:.2g}" for s in SPAD))
    print('  best R2 over cycles 1-5 (any source/measure):', max((r[8], r[1], r[2], r[3]) for r in rec if r[0] == BNAME[b] and r[1] != 'All cycles'))
